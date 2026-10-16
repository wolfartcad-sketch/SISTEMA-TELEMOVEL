import streamlit as st
import openpyxl
from datetime import datetime
import os
import re

# Lista oficial de serviços exata que aprovámos anteriormente
LISTA_SERVICOS = [
    "INSTALAÇÃO",
    "INSTALAÇÃO MAIS DE 300M",
    "TROCA DE ENDEREÇO",
    "TROCA DE ENDEREÇO MAIS DE 300M",
    "MANUTENÇÃO",
    "ROMPIMENTO",
    "MOVIMENTAR",
    "IMPRODUTIVA",
    "INSTALAÇÃO CAMERA WIFI AGREGADO",
    "INSTALAÇÃO CAMERA AGREGADO",
    "INSTALAÇÃO CAMERA AVULSO",
    "INSTALAÇÃO WIFI\\TV\\FONE\\PONTO EXTRA AVULSO",
    "INSTALAÇÃO WIFI\\TV\\FONE\\PONTO EXTRA AGREGADO",
    "BATIDA DE CTO 1/8",
    "BATIDA DE CTO 1/16",
    "RECOLHIMENTO EQUIPAMENTO ENDEREÇO ANTIGO",
    "INSTALAÇÃO FB",
    "INSTALAÇÃO MAIS DE 300M FB",
    "TROCA DE ENDEREÇO FB",
    "TROCA DE ENDEREÇO MAIS DE 300M FB",
    "MANUTENÇÃO FB",
    "ROMPIMENTO FB",
    "MOVIMENTAR FB",
    "IMPRODUTIVA FB",
    "INSTALAÇÃO CAMERA WIFI AGREGADO FB",
    "INSTALAÇÃO CAMERA AGREGADO FB",
    "INSTALAÇÃO CAMERA AVULSO FB",
    "INSTALAÇÃO WIFI\\TV\\FONE\\PONTO EXTRA AVULSO FB",
    "INSTALAÇÃO WIFI\\TV\\FONE\\PONTO EXTRA AGREGADO FB",
    "BATIDA DE CTO 1/8 FB",
    "BATIDA DE CTO 1/16 FB",
    "RECOLHIMENTO EQUIPAMENTO ENDEREÇO ANTIGO FB"
]

def traduzir_servico(servico_whatsapp):
    """Mapeia o texto do WhatsApp para a lista exata do Excel."""
    s = servico_whatsapp.upper()
    if "ATIVAÇÃO" in s or "INSTALAÇÃO" in s or "INTERNET" in s:
        if "300" in s and ("MAIS" in s or "~" in s and "300" in s):
            # Se for mais de 300m
            return "INSTALAÇÃO MAIS DE 300M"
        return "INSTALAÇÃO"
    elif "TROCA DE ENDEREÇO" in s:
        if "300" in s:
            return "TROCA DE ENDEREÇO MAIS DE 300M"
        return "TROCA DE ENDEREÇO"
    elif "MANUTENÇÃO" in s:
        return "MANUTENÇÃO"
    elif "ROMPIMENTO" in s:
        return "ROMPIMENTO"
    # Padrão por defeito caso não detete exatamente
    return "INSTALAÇÃO"

def processar_mensagem_whatsapp(texto):
    """Extrai os campos estruturados da mensagem do WhatsApp."""
    dados = {
        "cliente": "",
        "protocolo": "",
        "servico": "INSTALAÇÃO",
        "cidade": "",
        "metragem": "0",
        "equipamento": ""
    }
    
    linhas = texto.split("\n")
    for linha in linhas:
        linha_sup = linha.strip()
        if linha_sup.startswith("Cliente:"):
            dados["cliente"] = linha_sup.replace("Cliente:", "").strip()
        elif linha_sup.startswith("Protocolo:"):
            dados["protocolo"] = linha_sup.replace("Protocolo:", "").strip()
        elif linha_sup.startswith("Cidade:"):
            dados["cidade"] = linha_sup.replace("Cidade:", "").strip()
        elif linha_sup.startswith("Serviço:"):
            serv_bruto = linha_sup.replace("Serviço:", "").strip()
            dados["servico"] = traduzir_servico(serv_bruto)
        elif linha_sup.startswith("Cabo:") or linha_sup.startswith("Km:"):
            val = re.findall(r'\d+', linha_sup)
            if val and dados["metragem"] == "0":
                dados["metragem"] = val[0]
        elif linha_sup.startswith("Obs:"):
            obs = linha_sup.replace("Obs:", "").strip()
            dados["equipamento"] = obs

    return dados

def adicionar_ao_excel(caminho_arquivo, dados):
    if not os.path.exists(caminho_arquivo):
        return False, f"Ficheiro não encontrado: {caminho_arquivo}"

    wb = openpyxl.load_workbook(caminho_arquivo)
    ws = wb.active 

    proxima_linha = ws.max_row + 1
    for row in range(2, ws.max_row + 2):
        if ws.cell(row=row, column=2).value is None:
            proxima_linha = row
            break

    data_atual = datetime.now().strftime("%d/%m/%Y")

    ws.cell(row=proxima_linha, column=1, value=data_atual)
    ws.cell(row=proxima_linha, column=2, value=dados["cliente"])
    ws.cell(row=proxima_linha, column=3, value=dados["protocolo"])
    ws.cell(row=proxima_linha, column=4, value=dados["servico"])
    ws.cell(row=proxima_linha, column=5, value=dados["cidade"])
    ws.cell(row=proxima_linha, column=6, value=dados["metragem"])
    ws.cell(row=proxima_linha, column=8, value=dados["equipamento"])

    wb.save(caminho_arquivo)
    return True, proxima_linha

# --- INTERFACE STREAMLIT ---
st.title("Sistema de Atendimento - Automação Integrada ⚡")

st.subheader("📥 Importação Automática por Mensagem do WhatsApp")
msg_input = st.text_area("Cole a mensagem do WhatsApp aqui:", height=200)

caminho_excel = r"C:\Users\wolfa\OneDrive\Documentos\SERVIÇOS AUTOMATIZADOS MARCIO.xlsx"

if st.button("Processar Mensagem"):
    if msg_input.strip():
        dados_extraidos = processar_mensagem_whatsapp(msg_input)
        st.session_state["dados_processados"] = dados_extraidos
        st.success("Mensagem processada com sucesso!")
    else:
        st.warning("Por favor, cole uma mensagem válida.")

# Se já tiver dados processados na sessão, mostra o formulário de revisão e o botão de gravar
if "dados_processados" in st.session_state:
    d = st.session_state["dados_processados"]
    
    st.markdown("---")
    st.subheader("🔍 Rever Dados Extraídos antes de Enviar")
    
    col1, col2 = st.columns(2)
    with col1:
        cliente_edit = st.text_input("Cliente", value=d["cliente"])
        protocolo_edit = st.text_input("Protocolo", value=d["protocolo"])
        cidade_edit = st.text_input("Cidade", value=d["cidade"])
    with col2:
        # Seletor na lista oficial exata
        servico_edit = st.selectbox("Serviço", LISTA_SERVICOS, index=LISTA_SERVICOS.index(d["servico"]) if d["servico"] in LISTA_SERVICOS else 0)
        metragem_edit = st.text_input("Metragem (m)", value=d["metragem"])
        equipamento_edit = st.text_input("Equipamento / Obs", value=d["equipamento"])

    if st.button("🚀 Enviar para a Planilha Excel"):
        dados_finais = {
            "cliente": cliente_edit,
            "protocolo": protocolo_edit,
            "servico": servico_edit,
            "cidade": cidade_edit,
            "metragem": metragem_edit,
            "equipamento": equipamento_edit
        }
        
        sucesso, resultado = adicionar_ao_excel(caminho_excel, dados_finais)
        if sucesso:
            st.success(f"Atendimento de **{cliente_edit}** gravado com sucesso na linha **{resultado}** do Excel!")
            # Limpa o estado para o próximo atendimento
            del st.session_state["dados_processados"]
        else:
            st.error(resultado)