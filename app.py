import streamlit as st
import webbrowser
import time
import re
import pandas as pd
import openpyxl
import os
from datetime import datetime
import urllib.parse

IS_CLOUD = os.path.exists("/mount/src")

if not IS_CLOUD:
    import pyautogui
    import pyperclip
else:
    pyautogui = None
    pyperclip = None

# Configuração de página alterada para centrado e com limite elegante
st.set_page_config(page_title="Sistema de Atendimento - Assertec", layout="centered")

# CSS personalizado para limitar a largura máxima do container ao tamanho do título
st.markdown(
    """
    <style>
    .block-container {
        max-width: 900px;
        padding-top: 2rem;
        padding-right: 2rem;
        padding-left: 2rem;
        padding-bottom: 2rem;
    }
    </style>
    """,
    unsafe_allow_html=True
)

st.title("Sistema Automação Integrada FibraWer ⚡")

# --- LISTAS OFICIAIS DE APOIO ---
lista_cidades_assertec = [
    "Alto Feliz", "Bom Principio", "Caxias do Sul", "Feliz", "Harmonia", "Ivoti", 
    "Linha Nova", "Pareci Novo", "Presidente Lucena", "São José do Hortêncio", 
    "São Sebastião do Caí", "São Vendelino", "Tupandi", "Vale Real"
]

lista_servicos_assertec_global = [
    "Escolher...",
    "Atendimento Presencial - REPARO",
    "Atendimento Presencial - REPARO - ROMPIMENTO/LANÇAMENTO até 500 mt",
    "Atendimento presencial - RNC",
    "Visita improdutiva",
    "Desativações de serviços completa",
    "Desativações de serviços parcial / troca de endereço / migração de tecnologia",
    "Diária de atendimentos (sob demanda - dupla)",
    "Km rodado acima de 50km",
    "Ativação de Serviço - Acesso à Internet (0 ~ 300 mt)",
    "Ativação de Serviço - Acesso à Internet (301 ~ 500 mt)",
    "Ativação de Serviço - Agregado (câmeras)",
    "Ativação de Serviço - Agregado (câmeras wifi)",
    "Ativação de Serviço - Avulso (câmeras wifi)",
    "Ativação de Serviço - Agregado (wifi/tv/fone/ponto extra)",
    "Ativação de Serviço - Agregado Avulso (wifi/tv/fone/ponto extra)",
    "Batimento de caixa 1x16 p",
    "Batimento de caixa 1x8p",
    "Metragem excedente (acima de 500 mt)",
    "Troca de ponto interno (ponto principal/wifi/tv/fone/câmeras/ponto extra)",
    "Lançamento de fibra avulsa até f.o 6v",
    "Vistoria Técnica/ Viabilidade técnica",
    "Fusão em caixa",
    "Emenda de cabos rompidos na rede em infraestrutura própria",
    "Instalação de nova CTO",
    "Lançamento de cabo AS",
    "PLANTAO - duplo",
    "PLANTAO - simples",
    "PLANTAO - Atendimento Presencial - REPARO",
    "PLANTAO - Atendimento Presencial - REPARO - ROMPIMENTO/LANÇAMENTO até 500 mt",
    "CORPORATIVO - ATENDIMENTO PRESENCIAL - REPARO",
    "CORPORATIVO - ATENDIMENTO PRESENCIAL - REPARO - ROMPIMENTO/LANÇAMENTO ATÉ 500MT",
    "CORPORATIVO - ATENDIMENTO PRESENCIAL - RNC",
    "CORPORATIVO - ATIVAÇÃO DE SERVIÇO - ACESSO À INTERNET (0 ~ 300 MT)",
    "CORPORATIVO - ATIVAÇÃO DE SERVIÇO - ACESSO À INTERNET (301 ~ 500MT)",
    "CORPORATIVO - ATIVAÇÃO DE SERVIÇO - ACESSO À INTERNET (ACIMA DE 500MT)",
    "CORPORATIVO - ATIVAÇÃO DE SERVIÇO - ACESSO À INTERNET (RÁDIO)",
    "CORPORATIVO - ATIVAÇÃO DE SERVIÇO AGREGADO (CÂMERAS WIFI)",
    "CORPORATIVO - ATIVAÇÃO DE SERVIÇO AGREGADO (CÂMERAS)",
    "CORPORATIVO - ATIVAÇÃO DE SERVIÇO AGREGADO (WIFI/TV/FONE/PONTO EXTRA)",
    "CORPORATIVO - ATIVAÇÃO DE SERVIÇO AGREGADO AVULSO (CÂMERAS WIFI)",
    "CORPORATIVO - ATIVAÇÃO DE SERVIÇO AGREGADO AVULSO (WIFI/TV/FONE/PONTO EXTRA)",
    "CORPORATIVO - BATIMENTO DE CAIXA 1X16 PORTAS",
    "CORPORATIVO - BATIMENTO DE CAIXA 1X8 PORTAS",
    "CORPORATIVO - DESATIVAÇÃO DE SERVIÇO - COMPLETA",
    "CORPORATIVO - DESATIVAÇÃO DE SERVIÇO - PARCIAL",
    "CORPORATIVO - DIÁRIAS (SOB DEMANDA - DUPLA)",
    "CORPORATIVO - KM RODADO (ACIMA DE 50KM)",
    "CORPORATIVO - LANÇAMENTO DE FIBRA AVULSA ATÉ - FIBRA ATÉ 12V",
    "CORPORATIVO - METRAGEM EXCEDENTE (ACIMA DE 500MT)",
    "CORPORATIVO - TROCA DE PONTO INTERNO (PONTO PRINCIPAL/WIFI/TV/FONE/CÂMERAS/PONTO EXTRA)",
    "CORPORATIVO - VISITA IMPRODUTIVA",
    "CORPORATIVO - VISTORIA TÉCNICA / VIABILIDADE TÉCNICA"
]

lista_cidades_amigo = [
    "Escolher...", 
    "Alto Feliz", "Bom Principio", "Caxias do Sul", "Feliz", "Harmonia", "Ivoti", 
    "Linha Nova", "Pareci Novo", "Presidente Lucena", "São José do Hortêncio", 
    "São Sebastião do Caí", "São Vendelino", "Tupandi", "Vale Real"
]

lista_servicos_amigo_global = [
    "Escolher...", 
    "ATENDIMENTO PRESENCIAL - RNC", 
    "ATIVAÇÃO DE SERVIÇO - ACESSO À INTERNET (0 ~ 300 MT)", 
    "ATIVAÇÃO DE SERVIÇO - ACESSO À INTERNET (301 ~ 500MT)", 
    "ATIVAÇÃO DE SERVIÇO - ACESSO À INTERNET (ACIMA DE 500MT)", 
    "ATIVAÇÃO DE SERVIÇO - ACESSO À INTERNET (RÁDIO)", 
    "ATIVAÇÃO DE SERVIÇO AGREGADO (CÂMERAS WIFI)", 
    "ATIVAÇÃO DE SERVIÇO AGREGADO (Câmeras)", 
    "ATIVAÇÃO DE SERVIÇO AGREGADO (WIFI/TV/FONE/PONTO EXTRA)", 
    "ATIVAÇÃO DE SERVIÇO AGREGADO AVULSO (WIFI/TV/FONE/PONTO EXTRA)", 
    "ATIVAÇÃO DE SERVIÇO AGREGADO AVULSO (CÂMERAS WIFI)", 
    "BATIMENTO DE CAIXA 1X16 PORTAS", 
    "BATIMENTO DE CAIXA 1X8 PORTAS", 
    "DESATIVAÇÃO DE SERVIÇO - PARCIAL", 
    "DIÁRIAS (SOB DEMANDA - DUPLA)", 
    "LANÇAMENTO DE FIBRA AVULSA ATÉ - FIBRA ATÉ 12V", 
    "TROCA DE PONTO INTERNO (PONTO PRINCIPAL/WIFI/TV/FONE/CÂMERAS/PONTO EXTRA)", 
    "VISTORIA TÉCNICA / VIABILIDADE TÉCNICA"
]

# --- CRIAÇÃO DAS ABAS PRINCIPAIS ---
if IS_CLOUD:
    aba_gerador = st.container()
    aba_gestao = None
else:
    aba_gestao, aba_gerador = st.tabs(["💻 Gestão & Computador (Checklists & Excel)", "📱 Gerador de Mensagens (Telemóvel)"])

# ==========================================
# ABA 1: GESTÃO & COMPUTADOR
# ==========================================
if aba_gestao is not None:
    with aba_gestao:
        def limpar_campo():
            st.session_state["msg_input"] = ""
            st.session_state["mensagem_processada"] = False
            st.session_state["tec_msg"] = ""
            st.session_state["tec"] = ""
            st.session_state["cli"] = ""
            st.session_state["proto"] = ""
            st.session_state["cidade"] = ""
            st.session_state["servico"] = ""
            st.session_state["qtd"] = ""
            st.session_state["obs"] = ""

        if "mensagem_processada" not in st.session_state:
            st.session_state["mensagem_processada"] = False

        msg_texto = st.text_area(
            "Cole a mensagem do WhatsApp aqui (Ex: ASSERTEC \n MARCIO ...)...", 
            height=280, 
            key="msg_input"
        )

        if "formulario_escolhido" not in st.session_state:
            st.session_state["formulario_escolhido"] = "Assertec / Fibrawer"

        col_proc, col_limpar, col_espaco = st.columns([1.2, 1.2, 3])

        with col_proc:
            processar = st.button("Processar Mensagem")

        with col_limpar:
            st.button("Limpar Mensagem", on_click=limpar_campo)

        if processar:
            if not msg_texto.strip():
                st.error("Cole uma mensagem válida.")
            else:
                linhas = [l.strip() for l in msg_texto.strip().split("\n") if l.strip()]
                
                primeiro_texto = linhas[0].lower() if len(linhas) > 0 else ""
                if "amigo" in primeiro_texto or "tecpar" in primeiro_texto:
                    st.session_state["formulario_escolhido"] = "Brasil TecPar / Amigo"
                else:
                    st.session_state["formulario_escolhido"] = "Assertec / Fibrawer"

                if len(linhas) > 1:
                    st.session_state["tec_msg"] = linhas[1].upper()
                else:
                    st.session_state["tec_msg"] = ""

                def extrair(padrao, texto):
                    match = re.search(padrao, texto, re.IGNORECASE)
                    if match:
                        for val in match.groups():
                            if val is not None and val.strip():
                                return val.strip()
                    return ""
                
                st.session_state["tec"] = extrair(r"Técnico:\s*(.*)", msg_texto)
                st.session_state["cli"] = extrair(r"Cliente:\s*(.*)", msg_texto)
                st.session_state["proto"] = extrair(r"Protocolo:\s*(.*)", msg_texto)
                st.session_state["cidade"] = extrair(r"Cidade:\s*(.*)", msg_texto)
                
                raw_servico = extrair(r"Serviço:\s*(.*)", msg_texto)
                raw_servico = re.sub(r'ATATIVAÇÃO', 'ATIVAÇÃO', raw_servico, flags=re.IGNORECASE)
                st.session_state["servico"] = raw_servico
                
                qtd_extraida = extrair(r"Cabo[:]?\s*(.*)|Metragem[:]?\s*(.*)|Qtd[:]?\s*(.*)|Quantidade[:]?\s*(.*)|LPU[:]?\s*(.*)", msg_texto)
                st.session_state["qtd"] = qtd_extraida if qtd_extraida else "0"
                
                st.session_state["obs"] = extrair(r"Obs:\s*(.*)|Observação:\s*(.*)", msg_texto)
                st.session_state["mensagem_processada"] = True
                st.success(f"Mensagem processada! Formulário: **{st.session_state['formulario_escolhido']}** | Técnico Identificado: **{st.session_state['tec_msg']}**")

        def traduzir_para_excel(servico_wpp, metragem_str, cidade_atendimento):
            s = servico_wpp.upper()
            cidade_limpa = str(cidade_atendimento).strip().lower()
            
            # Lista das tuas cidades principais (comparações em minúsculas)
            cidades_principais = [
                "ivoti", 
                "são josé do hortêncio", 
                "sao jose do hortencio", 
                "linha nova", 
                "são sebastião do caí", 
                "sao sebastiao do cai", 
                "presidente lucena"
            ]
            
            # Verifica se a cidade é de fora da base
            cidade_fora_da_base = cidade_limpa and cidade_limpa not in cidades_principais
            
            # Verifica se é um serviço de plantão específico considerado fora da base
            eh_plantao_fb = "PLANTAO" in s and ("REPARO" in s)
            
            # Define se precisa do sufixo FB no final
            sufixo_fb = " FB" if (cidade_fora_da_base or eh_plantao_fb) else ""

            try:
                m = int(re.sub(r'\D', '', str(metragem_str)))
            except:
                m = 0
            
            if "REPARO" in s:
                return f"MANUTENÇAO{sufixo_fb}"
            elif "ATIVAÇÃO" in s or "INSTALAÇÃO" in s or "INTERNET" in s:
                if m > 300:
                    return f"INSTALAÇAO MAIS DE 300M{sufixo_fb}"
                return f"INSTALAÇAO{sufixo_fb}"
            elif "TROCA DE ENDEREÇO" in s:
                if m > 300:
                    return f"TROCA DE ENDEREÇO MAIS DE 300M{sufixo_fb}"
                return f"TROCA DE ENDEREÇO{sufixo_fb}"
            elif "MANUTENÇÃO" in s or "MANUTENCAO" in s:
                return f"MANUTENÇAO{sufixo_fb}"
            elif "ROMPIMENTO" in s:
                return f"ROMPIMENTO{sufixo_fb}"
            elif "IMPRODUTIVA" in s or "VISITA IMPRODUTIVA" in s:
                return f"IMPRODUTIVA{sufixo_fb}"
            elif "BATIDA" in s:
                if "1/16" in s or "1X16" in s:
                    return f"BATIDA DE CTO 1/16{sufixo_fb}"
                return f"BATIDA DE CTO 1/8{sufixo_fb}"
            
            return f"INSTALAÇAO{sufixo_fb}"

        def obter_caminho_excel_tecnico(tec_nome):
            nome_limpo = str(tec_nome).strip().upper()
            if not nome_limpo:
                nome_limpo = "MARCIO"
            return rf"C:\Users\wolfa\OneDrive\Documentos\SERVIÇOS AUTOMATIZADOS {nome_limpo}.xlsx"

        def adicionar_ao_excel(caminho_arquivo, dados):
            if not os.path.exists(caminho_arquivo):
                return False, f"Ficheiro não encontrado: {caminho_arquivo}"
            try:
                wb = openpyxl.load_workbook(caminho_arquivo)
                if "LANÇAMENTOS" in wb.sheetnames:
                    ws = wb["LANÇAMENTOS"]
                else:
                    ws = wb.active 

                proxima_linha = ws.max_row + 1
                for row in range(2, ws.max_row + 2):
                    if ws.cell(row=row, column=2).value is None:
                        proxima_linha = row
                        break

                data_atual = datetime.now().strftime("%d/%m/%Y")

                ws.cell(row=proxima_linha, column=1, value=data_atual)
                ws.cell(row=proxima_linha, column=2, value=dados["cliente"])
                ws.cell(row=proxima_linha, column=3, value=dados["protocolo"])
                ws.cell(row=proxima_linha, column=4, value=dados["servico"])
                ws.cell(row=proxima_linha, column=5, value=dados["cidade"])
                ws.cell(row=proxima_linha, column=6, value=dados["metragem"])
                ws.cell(row=proxima_linha, column=8, value=dados["equipamento"])

                wb.save(caminho_arquivo)
                return True, proxima_linha
            except Exception as e:
                return False, str(e)

        def conferir_duplicidade_servico(caminho_arquivo, protocolo, servico_desejado):
            if not os.path.exists(caminho_arquivo):
                return False, 0
            try:
                wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
                if "LANÇAMENTOS" in wb.sheetnames:
                    ws = wb["LANÇAMENTOS"]
                else:
                    ws = wb.active
                    
                proto_limpo = str(protocolo).strip().lower()
                serv_limpo = str(servico_desejado).strip().lower()
                
                for row in range(2, ws.max_row + 1):
                    col_proto = str(ws.cell(row=row, column=3).value or "").strip().lower()
                    col_serv = str(ws.cell(row=row, column=4).value or "").strip().lower()
                    
                    if col_proto == proto_limpo and col_serv == serv_limpo:
                        return True, row
                return False, 0
            except Exception as e:
                return False, 0

        def conferir_na_planilha_local(caminho_arquivo, termo):
            if not os.path.exists(caminho_arquivo):
                return False, "Ficheiro Excel não encontrado."
            try:
                wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
                if "LANÇAMENTOS" in wb.sheetnames:
                    ws = wb["LANÇAMENTOS"]
                else:
                    ws = wb.active
                    
                termo_limpo = termo.strip().lower()
                for row in range(2, ws.max_row + 1):
                    cliente = str(ws.cell(row=row, column=2).value or "").lower()
                    protocolo = str(ws.cell(row=row, column=3).value or "").lower()
                    if termo_limpo in cliente or termo_limpo in protocolo:
                        return True, row
                return False, 0
            except Exception as e:
                return False, str(e)

        if st.session_state["mensagem_processada"]:
            st.markdown("---")
            formulario_escolhido = st.session_state["formulario_escolhido"]
            
            if formulario_escolhido == "Assertec / Fibrawer":
                btn_executar = st.button("🚀 Executar Assertec", key="btn_exec_ass_topo")
                if btn_executar:
                    tecnico_val = st.session_state.get("tec", "")
                    cliente_val = st.session_state.get("cli", "")
                    protocolo_val = st.session_state.get("proto", "")
                    cidade_importada = st.session_state.get("cidade", "")
                    servico_importado = st.session_state.get("servico", "")
                    quantidade_val = st.session_state.get("qtd", "") if st.session_state.get("qtd", "").strip() else "1"
                    observacao_val = st.session_state.get("obs", "")
                    
                    cidade_idx = 0
                    for idx, c in enumerate(lista_cidades_assertec):
                        if c.lower() in cidade_importada.lower() and cidade_importada.strip():
                            cidade_idx = idx
                            break
                    
                    servico_idx = 0
                    serv_imp_clean = servico_importado.strip().lower()
                    if serv_imp_clean:
                        for idx, s in enumerate(lista_servicos_assertec_global):
                            if s.strip().lower() == serv_imp_clean:
                                servico_idx = idx
                                break
                        if servico_idx == 0:
                            for idx, s in enumerate(lista_servicos_assertec_global):
                                if serv_imp_clean in s.strip().lower() or s.strip().lower() in serv_imp_clean:
                                    servico_idx = idx
                                    break

                    if not tecnico_val.strip() or not cliente_val.strip() or not protocolo_val.strip():
                        st.error("Preencha o técnico, o cliente e o protocolo.")
                    else:
                        st.warning("⚡ A executar automação Assertec... Não mexa no rato.")
                        webbrowser.open("https://forms.gle/2MENn9xuKjj7pKVm6")
                        time.sleep(2.0)
                        
                        for _ in range(3):
                            pyautogui.press('tab')
                            time.sleep(0.04)
                        pyautogui.press('tab')
                        pyautogui.press('tab')
                        pyautogui.press('tab')
                        pyautogui.press('enter')
                        time.sleep(0.2)
                        pyautogui.press('enter')
                        time.sleep(0.08)
                        
                        pyautogui.press('tab')
                        time.sleep(0.04)
                        pyautogui.press('enter')
                        time.sleep(0.25)
                        pyautogui.write(tecnico_val)
                        time.sleep(0.1)
                        pyautogui.press('enter')
                        time.sleep(0.1)
                        
                        pyautogui.press('tab')
                        pyautogui.press('tab')
                        pyautogui.press('space')
                        time.sleep(0.1)
                        
                        for _ in range(3):
                            pyautogui.press('tab')
                            time.sleep(0.04)
                        
                        pyperclip.copy(str(cliente_val))
                        pyautogui.hotkey('ctrl', 'v')
                        time.sleep(0.08)
                        
                        pyautogui.press('tab')
                        pyperclip.copy(str(protocolo_val))
                        pyautogui.hotkey('ctrl', 'v')
                        time.sleep(0.08)
                        
                        pyautogui.press('tab')
                        for _ in range(cidade_idx):
                            pyautogui.press('tab')
                            time.sleep(0.04)
                        
                        pyautogui.press('space')
                        time.sleep(0.1)
                        
                        passos_restantes = len(lista_cidades_assertec) - cidade_idx
                        for _ in range(passos_restantes):
                            pyautogui.press('tab')
                            time.sleep(0.04)
                        
                        pyautogui.press('enter')
                        time.sleep(0.3)
                        
                        pyautogui.press('home')
                        time.sleep(0.1)
                        for _ in range(servico_idx):
                            pyautogui.press('down')
                            time.sleep(0.03)
                        
                        pyautogui.press('enter')
                        time.sleep(0.2)
                        
                        pyautogui.press('tab')
                        time.sleep(0.1)
                        
                        pyautogui.press('tab')
                        time.sleep(0.1)
                        
                        pyperclip.copy(str(quantidade_val))
                        pyautogui.hotkey('ctrl', 'v')
                        time.sleep(0.1)
                        
                        pyautogui.press('tab')
                        time.sleep(0.1)
                        
                        if observacao_val.strip():
                            pyperclip.copy(str(observacao_val))
                            pyautogui.hotkey('ctrl', 'v')
                            time.sleep(0.1)
                        
                        st.success("Automação Assertec concluída com sucesso! 🚀")

            elif formulario_escolhido == "Brasil TecPar / Amigo":
                btn_executar_amigo = st.button("🚀 Executar Amigo", key="btn_exec_amig_topo")
                if btn_executar_amigo:
                    tecnico_amigo_val = "ANDERSON WOLFART MOLLERKE"
                    cidade_importada_amigo = st.session_state.get("cidade", "")
                    servico_importado_amigo = st.session_state.get("servico", "")
                    cliente_val = st.session_state.get("cli", "")
                    protocolo_val = st.session_state.get("proto", "")
                    quantidade_cabo_val = st.session_state.get("qtd", "") if st.session_state.get("qtd", "").strip() else "0"
                    observacao_val = st.session_state.get("obs", "")
                    
                    cidade_idx_amigo = 0
                    for idx, c in enumerate(lista_cidades_amigo):
                        if c.lower() in cidade_importada_amigo.lower() and cidade_importada_amigo.strip():
                            cidade_idx_amigo = idx
                            break
                    
                    servico_idx_amigo = 0
                    serv_imp_clean_amigo = servico_importado_amigo.strip().lower()
                    if serv_imp_clean_amigo:
                        for idx, s in enumerate(lista_servicos_amigo_global):
                            if s.strip().lower() == serv_imp_clean_amigo:
                                servico_idx_amigo = idx
                                break
                        if servico_idx_amigo == 0:
                            for idx, s in enumerate(lista_servicos_amigo_global):
                                if serv_imp_clean_amigo in s.strip().lower():
                                    servico_idx_amigo = idx
                                    break

                    st.warning("⚡ A executar automação Brasil TecPar... Não mexa no rato nem no teclado.")
                    webbrowser.open("https://docs.google.com/forms/d/e/1FAIpQLSfDT9GbrgCv5vcOi2fvZoeuPITItTWIQFWmS2vO_NY7gMEvQA/viewform")
                    time.sleep(2.5)
                    
                    for _ in range(3):
                        pyautogui.press('tab')
                        time.sleep(0.04)
                    
                    pyautogui.press('enter')
                    time.sleep(0.3)
                    pyautogui.write("ASSERTEC SERVICOS E TECNOLOGIA LTDA")
                    time.sleep(0.15)
                    pyautogui.press('enter')
                    time.sleep(0.15)
                    
                    pyautogui.press('tab')
                    time.sleep(0.05)
                    pyautogui.press('enter')
                    time.sleep(0.3)
                    pyautogui.write(tecnico_amigo_val)
                    time.sleep(0.15)
                    pyautogui.press('enter')
                    time.sleep(0.15)
                    
                    pyautogui.press('tab')
                    time.sleep(0.05)
                    pyautogui.press('enter')
                    time.sleep(1.2)
                    
                    for _ in range(3):
                        pyautogui.press('tab')
                        time.sleep(0.04)
                        
                    pyperclip.copy(str(protocolo_val))
                    pyautogui.hotkey('ctrl', 'v')
                    time.sleep(0.15)
                    
                    pyautogui.press('tab')
                    time.sleep(0.05)
                    pyautogui.press('enter')
                    time.sleep(0.3)
                    pyautogui.write("AMIGO - VAREJO")
                    time.sleep(0.15)
                    pyautogui.press('enter')
                    time.sleep(0.15)
                    
                    pyautogui.press('tab')
                    time.sleep(0.05)
                    pyperclip.copy(str(cliente_val))
                    pyautogui.hotkey('ctrl', 'v')
                    time.sleep(0.15)
                    
                    pyautogui.press('tab')
                    time.sleep(0.05)
                    pyautogui.press('enter')
                    time.sleep(0.3)
                    
                    pyautogui.press('home')
                    time.sleep(0.1)
                    for _ in range(cidade_idx_amigo):
                        pyautogui.press('down')
                        time.sleep(0.03)
                    pyautogui.press('enter')
                    time.sleep(0.15)
                    
                    pyautogui.press('tab')
                    time.sleep(0.1)
                    pyautogui.press('tab')
                    time.sleep(0.1)
                    pyperclip.copy("0")
                    pyautogui.hotkey('ctrl', 'v')
                    time.sleep(0.15)
                    
                    pyautogui.press('tab')
                    time.sleep(0.05)
                    pyautogui.press('tab')
                    time.sleep(0.05)
                    pyautogui.press('enter')
                    time.sleep(1.2)
                    
                    for _ in range(3):
                        pyautogui.press('tab')
                        time.sleep(0.04)
                        
                    pyautogui.press('enter')
                    time.sleep(0.3)
                    
                    pyautogui.press('home')
                    time.sleep(0.1)
                    for _ in range(servico_idx_amigo):
                        pyautogui.press('down')
                        time.sleep(0.03)
                    
                    pyautogui.press('enter')
                    time.sleep(0.3)
                    pyautogui.press('tab')
                    time.sleep(0.05)
                    pyautogui.press('tab')
                    time.sleep(0.05)
                    pyautogui.press('tab')
                    time.sleep(0.05)
                    pyautogui.press('enter')
                    time.sleep(1.2)
                    
                    servico_atual_upper = servico_importado_amigo.upper()
                    pula_pagina_cabo = "BATIMENTO" in servico_atual_upper or "DESATIVAÇÃO" in servico_atual_upper or "DIÁRIAS" in servico_atual_upper or "RNC" in servico_atual_upper or "VISTORIA" in servico_atual_upper
                    
                    if not pula_pagina_cabo:
                        for _ in range(3):
                            pyautogui.press('tab')
                            time.sleep(0.04)
                            
                        pyperclip.copy(str(quantidade_cabo_val))
                        pyautogui.hotkey('ctrl', 'v')
                        time.sleep(0.15)
                        
                        pyautogui.press('tab')
                        time.sleep(0.05)
                        pyautogui.press('tab')
                        time.sleep(0.05)
                        pyautogui.press('enter')
                        time.sleep(1.2)
                    
                    for _ in range(3):
                        pyautogui.press('tab')
                        time.sleep(0.04)
                        
                    pyperclip.copy("1")
                    pyautogui.hotkey('ctrl', 'v')
                    time.sleep(0.15)
                    
                    pyautogui.press('tab')
                    time.sleep(0.05)
                    pyautogui.press('tab')
                    time.sleep(0.05)
                    pyautogui.press('enter')
                    time.sleep(1.2)
                    
                    for _ in range(3):
                        pyautogui.press('tab')
                        time.sleep(0.04)
                        
                    if observacao_val.strip():
                        pyperclip.copy(str(observacao_val))
                        pyautogui.hotkey('ctrl', 'v')
                        time.sleep(0.15)
                    
                    st.success("Automação Brasil TecPar concluída até às Observações! 🚀")

            st.markdown("")
            col_lancar, col_conferir_plan, col_vazio_botoes = st.columns([1.5, 1.5, 2])
            
            tec_msg_atual = st.session_state.get("tec_msg", "MARCIO")
            caminho_excel = obter_caminho_excel_tecnico(tec_msg_atual)
            
            with col_lancar:
                if st.button("🚀 Lançar na Planilha Excel"):
                    cli_v = st.session_state.get("cli", "")
                    proto_v = st.session_state.get("proto", "")
                    serv_v = st.session_state.get("servico", "")
                    cid_v = st.session_state.get("cidade", "")
                    qtd_v = st.session_state.get("qtd", "0")
                    obs_v = st.session_state.get("obs", "")
                    
                    servico_convertido = traduzir_para_excel(serv_v, qtd_v, cid_v)
                    ja_existe, linha_existente = conferir_duplicidade_servico(caminho_excel, proto_v, servico_convertido)
                    
                    if ja_existe:
                        st.warning(f"⚠️ Atenção: O serviço **'{servico_convertido}'** para o protocolo **{proto_v}** já foi lançado na linha **{linha_existente}** do ficheiro de **{tec_msg_atual}**!")
                    else:
                        dados_registo = {
                            "cliente": cli_v,
                            "protocolo": proto_v,
                            "servico": servico_convertido,
                            "cidade": cid_v,
                            "metragem": re.sub(r'\D', '', str(qtd_v)) or "0",
                            "equipamento": obs_v
                        }
                        sucesso, res = adicionar_ao_excel(caminho_excel, dados_registo)
                        if sucesso:
                            st.success(f"✅ Atendimento de **{cli_v}** adicionado na linha **{res}** da planilha de **{tec_msg_atual}** (aba LANÇAMENTOS)!")
                        else:
                            st.error(f"Erro ao gravar no ficheiro de {tec_msg_atual}: {res}")
                        
            with col_conferir_plan:
                if st.button("🔍 Conferir se Já Existe"):
                    proto_v = st.session_state.get("proto", "")
                    cli_v = st.session_state.get("cli", "")
                    termo_busca = proto_v if proto_v.strip() else cli_v
                    
                    if not termo_busca.strip():
                        st.warning("Nenhum protocolo ou cliente detetado para conferir.")
                    else:
                        encontrado, linha_ou_erro = conferir_na_planilha_local(caminho_excel, termo_busca)
                        if encontrado:
                            st.success(f"✅ Encontrado na planilha de **{tec_msg_atual}** na linha **{linha_ou_erro}**!")
                        else:
                            st.info(f"ℹ️ O termo **'{termo_busca}'** ainda não foi encontrado na planilha de **{tec_msg_atual}**.")

        st.markdown("---")
        st.subheader("🔍 Consulta Global de Cliente / Protocolo")
        termo_pesquisa_global = st.text_input("Consultar por Nome do Cliente ou Protocolo:", value=st.session_state.get("cli", ""), key="pesquisa_global_topo")

        col_conf_ass_topo, col_conf_amig_topo, col_espaco_topo = st.columns([1.2, 1.2, 3])
        with col_conf_ass_topo:
            btn_conf_ass_topo = st.button("🔍 Conferir Assertec", key="btn_conf_ass_topo_key")
        with col_conf_amig_topo:
            btn_conf_amig_topo = st.button("🔍 Conferir Amigo", key="btn_conf_amig_topo_key")

        if btn_conf_ass_topo:
            if not termo_pesquisa_global.strip():
                st.warning("Insira o nome de um cliente ou protocolo na caixa de pesquisa.")
            else:
                with st.spinner("A consultar a planilha da Assertec..."):
                    try:
                        url_planilha = "https://docs.google.com/spreadsheets/d/1MMQ8PAjr2GQfWxPnfYNfHSiiu9AUdE0pvfJikSGQgqU/export?format=csv"
                        df = pd.read_csv(url_planilha)
                        encontrado = False
                        termo_busca = termo_pesquisa_global.strip().lower()
                        for col in df.columns:
                            matches = df[df[col].astype(str).str.lower().str.contains(termo_busca, na=False)]
                            if not matches.empty:
                                encontrado = True
                                break
                        if encontrado:
                            st.success(f"✅ O registo com **'{termo_pesquisa_global}'** foi encontrado na planilha da Assertec!")
                            st.dataframe(matches)
                        else:
                            st.info(f"ℹ️ O termo **'{termo_pesquisa_global}'** ainda não consta na planilha da Assertec.")
                    except Exception:
                        st.warning("Não foi possível ler a planilha automaticamente.")

        if btn_conf_amig_topo:
            if not termo_pesquisa_global.strip():
                st.warning("Insira o nome do cliente ou protocolo na caixa de pesquisa.")
            else:
                with st.spinner("A abrir a planilha da Amigo..."):
                    webbrowser.open("https://docs.google.com/spreadsheets/d/1jESKPwkNUo0ocIPOUCzKKeq8wav57Rdd_K9MF3WBSDk/edit?gid=1573921361#gid=1573921361")
                    time.sleep(3.5)
                    pyautogui.hotkey('ctrl', 'f')
                    time.sleep(0.4)
                    pyperclip.copy(str(termo_pesquisa_global))
                    pyautogui.hotkey('ctrl', 'v')
                    time.sleep(0.3)
                    pyautogui.press('enter')
                st.success(f"✅ Planilha Amigo aberta! Procurou por **'{termo_pesquisa_global}'**.")

        st.markdown("---")
        
        if st.session_state.get("mensagem_processada", False):
            formulario_escolhido = st.session_state["formulario_escolhido"]
            st.info(f"📋 Formulário ativo no momento: **{formulario_escolhido}** | Técnico da Mensagem: **{st.session_state.get('tec_msg', 'N/D')}**")

            st.subheader(f"📋 Relatório de Conferência - {formulario_escolhido}")
            
            col_d1, col_d2 = st.columns(2)
            with col_d1:
                st.markdown(f"**Cliente:** {st.session_state.get('cli', 'N/D')}")
                st.markdown(f"**Protocolo:** {st.session_state.get('proto', 'N/D')}")
                st.markdown(f"**Cidade:** {st.session_state.get('cidade', 'N/D')}")
            with col_d2:
                st.markdown(f"**Serviço:** {st.session_state.get('servico', 'N/D')}")
                st.markdown(f"**Técnico:** {st.session_state.get('tec', 'N/D')}")
                st.markdown(f"**Quantidade / Metragem:** {st.session_state.get('qtd', '0')}")
            
            st.markdown(f"**Observações Finais:** {st.session_state.get('obs', 'Nenhuma observação.')}")
        else:
            st.info("ℹ️ Cole e processe uma mensagem do WhatsApp na caixa acima para visualizar o relatório de conferência.")

# ==========================================
# ==========================================
# ABA 2: GERADOR DE MENSAGENS (Telemóvel)
# ==========================================
with aba_gerador:
    st.markdown("""
        <style>
        input[aria-label*="*"], select[aria-label*="*"], div[data-baseweb="select"] > div:has(~ input[aria-label*="*"]) {
            border: 2px solid #ff4b4b !important;
            background-color: rgba(255, 75, 75, 0.03) !important;
        }
        label:has(span:contains("*")), p:has(span:contains("*")) {
            color: #d32f2f !important;
            font-weight: 700 !important;
        }
        </style>
    """, unsafe_allow_html=True)

    st.subheader("📱 Gerador de Mensagem Padronizada para WhatsApp")
    st.markdown("Selecione a empresa e preencha todos os campos obrigatórios para liberar a geração da mensagem.")

    # Empresa inicia em "Escolher..."
    opcoes_empresa = ["Escolher...", "ASSERTEC", "AMIGO"]
    empresa_gen = st.selectbox("Empresa / Tipo *", opcoes_empresa, key="gen_empresa")

    # Variáveis de controlo para validação
    campos_preenchidos = False

    if empresa_gen != "Escolher...":
        if empresa_gen == "ASSERTEC":
            st.error("🚨 **Atenção: Preencha todos os campos obrigatórios em destaque vermelho abaixo.**")
            
            tec_opcoes = ["Escolher...", "ANDERSON", "MARCIO"]
            tecnico_selecionado = st.selectbox("Técnico Responsável *", tec_opcoes, key="gen_tecnico_assertec")
            
            cidades_opcoes = ["Escolher..."] + lista_cidades_assertec
            cidade_gen = st.selectbox("Cidade *", cidades_opcoes, key="gen_cidade_assertec")
            
            servicos_opcoes = lista_servicos_assertec_global # Já tem "Escolher..." no topo
            servico_gen = st.selectbox("Serviço Realizado *", servicos_opcoes, key="gen_servico_assertec")
            
            cliente_gen = st.text_input("Nome do Cliente *", placeholder="EX: ELISETE ANTONIA DOS SANTOS", key="gen_cliente_assertec")
            protocolo_gen = st.text_input("Número do Protocolo *", placeholder="EX: 114619", key="gen_protocolo_assertec")
            metragem_gen = st.number_input("Metragem / Quantidade", min_value=0, value=1, step=1, key="gen_metragem_assertec")
            
            # Validação para Assertec
            if (tecnico_selecionado != "Escolher..." and 
                cidade_gen != "Escolher..." and 
                servico_gen != "Escolher..." and 
                cliente_gen.strip() != "" and 
                protocolo_gen.strip() != ""):
                campos_preenchidos = True

        else: # AMIGO
            st.error("🚨 **Atenção: Preencha todos os campos obrigatórios em destaque vermelho abaixo.**")
            
            tec_opcoes = ["Escolher...", "ANDERSON", "MARCIO"]
            tecnico_selecionado = st.selectbox("Técnico Responsável *", tec_opcoes, key="gen_tecnico_amigo")
            
            cidades_opcoes = ["Escolher..."] + lista_cidades_amigo
            cidade_gen = st.selectbox("Cidade *", cidades_opcoes, key="gen_cidade_amigo")
            
            servicos_opcoes = lista_servicos_amigo_global # Já tem "Escolher..." no topo
            servico_gen = st.selectbox("Serviço Realizado *", servicos_opcoes, key="gen_servico_amigo")
            
            cliente_gen = st.text_input("Nome do Cliente *", placeholder="EX: ELISETE ANTONIA DOS SANTOS", key="gen_cliente_amigo")
            protocolo_gen = st.text_input("Número do Protocolo *", placeholder="EX: 114619", key="gen_protocolo_amigo")
            
            lpu_gen = st.number_input("LPU", min_value=0, value=1, step=1, key="gen_lpu_amigo")
            metragem_amigo_gen = st.number_input("Metragem / Cabo", min_value=0, value=0, step=1, key="gen_metragem_amigo")
            km_rodado_gen = 0
            st.text_input("Quilômetro Rodado (Km)", value="0", disabled=True, key="gen_km_amigo_fixo")
            
            # Validação para Amigo
            if (tecnico_selecionado != "Escolher..." and 
                cidade_gen != "Escolher..." and 
                servico_gen != "Escolher..." and 
                cliente_gen.strip() != "" and 
                protocolo_gen.strip() != ""):
                campos_preenchidos = True

        st.markdown("---")
        obs_gen = st.text_area("Observações (Opcional)", placeholder="EX: Refeita configurações nos equipamentos...", key="gen_obs")

        # Configuração do nome completo do técnico para a mensagem
        if tecnico_selecionado == "ANDERSON":
            tecnico_gen = "ANDERSON"
            nome_tec_completo = "ANDERSON WOLFART MOLLERKE" if empresa_gen == "AMIGO" else "Anderson"
        else:
            tecnico_gen = "MARCIO"
            nome_tec_completo = "ANDERSON WOLFART MOLLERKE" if empresa_gen == "AMIGO" else "Márcio"

        # O botão só executa e gera se todos os campos obrigatórios estiverem preenchidos
        if campos_preenchidos:
            if st.button("✨ Gerar Mensagem Pronta"):
                if empresa_gen == "AMIGO":
                    mensagem_gerada = f"""{empresa_gen}
{tecnico_gen}
Técnico: {nome_tec_completo}
Cliente: {cliente_gen.upper()}
Protocolo: {protocolo_gen}
Cidade: {cidade_gen}
Serviço: {servico_gen}
LPU: {lpu_gen}
Metragem: {metragem_amigo_gen}
Km Rodado: {km_rodado_gen}
Obs: {obs_gen}"""
                else:
                    mensagem_gerada = f"""{empresa_gen}
{tecnico_gen}
Técnico: {nome_tec_completo}
Cliente: {cliente_gen.upper()}
Protocolo: {protocolo_gen}
Cidade: {cidade_gen}
Serviço: {servico_gen}
Metragem: {metragem_gen}
Obs: {obs_gen}"""
                st.session_state["mensagem_gerada_final"] = mensagem_gerada
        else:
            st.warning("⚠️ Por favor, preencha todos os campos obrigatórios destacados em vermelho para liberar a geração da mensagem.")

    if "mensagem_gerada_final" in st.session_state and st.session_state["mensagem_gerada_final"]:
        st.markdown("### 📋 Mensagem Gerada:")
        st.code(st.session_state["mensagem_gerada_final"], language="text")
        
        mensagem_url = urllib.parse.quote(st.session_state["mensagem_gerada_final"])
        link_whatsapp = f"https://wa.me/?text={mensagem_url}"
        st.link_button("📲 Enviar diretamente para o WhatsApp", link_whatsapp)